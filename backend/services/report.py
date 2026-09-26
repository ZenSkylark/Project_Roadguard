from pathlib import Path
from docx import Document
from docx.shared import Inches
from ..config import settings
from ..models import Violation

def generate_report(v: Violation) -> str:
    doc = Document()
    doc.add_heading("Roadguard Violation Report", 0)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    table.rows[0].cells[0].text = "Field"
    table.rows[0].cells[1].text = "Value"
    rows = [("Event ID", v.event_id), ("Violation", v.violation_type),
            ("Captured At", str(v.captured_at)),
            ("AI Confidence", f"{v.confidence:.2f}"),
            ("License Plate", v.plate_text or "NOT FILLED"),
            ("Plate Source", v.plate_source or "-"),
            ("Plate Confidence", f"{v.plate_confidence:.2f}" if v.plate_confidence else "-"),
            ("Status", v.status)]
    for k, val in rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = k, str(val)
    if v.image_path and Path(v.image_path).exists():
        try:
            doc.add_picture(v.image_path, width=Inches(5.5))
        except Exception:
            doc.add_paragraph("[Evidence image unavailable or corrupt]")
    out = Path(settings.STORAGE_ROOT) / "reports" / f"report_{v.id}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return str(out)