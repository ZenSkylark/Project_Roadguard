import json
from pathlib import Path
from ..config import settings

import tempfile
from docx import Document


def _get_settings_file() -> Path:
    """Compute settings path dynamically so monkeypatched STORAGE_ROOT is respected."""
    return Path(settings.STORAGE_ROOT) / "system_settings.json"


def load_ui_settings() -> dict:
    sf = _get_settings_file()
    if sf.exists():
        try:
            return json.loads(sf.read_text())
        except Exception:
            pass
    return {"template_dir": "./templates", "default_template": None}


def save_ui_settings(data: dict):
    sf = _get_settings_file()
    sf.parent.mkdir(parents=True, exist_ok=True)
    current = load_ui_settings()
    current.update(data)
    sf.write_text(json.dumps(current, indent=2))


def list_templates() -> list:
    s = load_ui_settings()
    tdir = Path(s.get("template_dir", "./templates"))
    if not tdir.exists():
        return []
    return sorted([f.name for f in tdir.glob("*.docx")])

def fill_template(template_name: str, data: dict) -> str:
    """Fill a .docx template with mail-merge data and return path to the filled file."""
    s = load_ui_settings()
    tdir = Path(s.get("template_dir", "./templates"))
    tpath = tdir / template_name

    if not tpath.exists():
        raise FileNotFoundError(f"Template not found: {tpath}")

    doc = Document(str(tpath))

    # Replace placeholders in paragraphs
    for paragraph in doc.paragraphs:
        for key, value in data.items():
            placeholder = "{{" + key + "}}"
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, str(value or ""))

    # Replace placeholders in tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for key, value in data.items():
                        placeholder = "{{" + key + "}}"
                        if placeholder in paragraph.text:
                            paragraph.text = paragraph.text.replace(placeholder, str(value or ""))

    # Save to temp file
    out = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(out.name)
    out.close()
    return out.name