import os
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from sqlalchemy.orm import Session
from ..audit import audit
from ..auth import require_position, get_current_user
from ..database import get_db
from ..models import Violation, User
from ..schemas import PlateIn
from ..services.report import generate_report
from ..clock import utcnow
from pathlib import Path
from ..config import settings



router = APIRouter(prefix="/api/violations", tags=["violations"])
ANY = ("administrator", "officer", "viewer")

def _get(db: Session, vid: int) -> Violation:
    v = db.get(Violation, vid)
    if not v or v.status == "deleted":
        raise HTTPException(404, "Violation not found")
    return v

@router.get("")
def list_violations(status: str | None = None, db: Session = Depends(get_db),
                    user=Depends(require_position(*ANY))):
    q = db.query(Violation)
    if status:
        q = q.filter(Violation.status == status)
    else:
        q = q.filter(Violation.status != "deleted")
    return [{"id": v.id, "event_id": v.event_id, "violation_type": v.violation_type,
             "confidence": v.confidence, "captured_at": v.captured_at,
             "plate_text": v.plate_text, "plate_source": v.plate_source,
             "status": v.status}
            for v in q.order_by(Violation.id.desc()).all()]

@router.get("/{vid}")
def get_violation(vid: int, db: Session = Depends(get_db),
                  user=Depends(require_position(*ANY))):
    v = _get(db, vid)
    return {
        "id": v.id,
        "event_id": v.event_id,
        "violation_type": v.violation_type,
        "confidence": v.confidence,
        "captured_at": v.captured_at,
        "plate_text": v.plate_text,
        "plate_source": v.plate_source,
        "status": v.status,
        "image_path": v.image_path,
        "uploaded_by": v.uploaded_by
    }

@router.patch("/{vid}/status")
def update_status(vid: int, body: dict, db: Session = Depends(get_db),
                  user=Depends(require_position("officer", "administrator"))):
    v = _get(db, vid)
    new_status = body.get("status")
    if new_status not in ("pending", "approved", "finalized", "rejected"):
        raise HTTPException(422, "Invalid status")
    v.status = new_status
    db.commit()
    audit(db, user.id, "violation_status_changed", f"violation={vid} status={new_status}")
    return {"id": v.id, "status": v.status}

@router.patch("/{vid}/plate")
def manual_plate(vid: int, body: PlateIn, db: Session = Depends(get_db),
                 user=Depends(require_position("officer", "administrator"))):
    v = _get(db, vid)
    v.plate_text = body.plate_text
    v.plate_source = "manual"
    v.report_path = generate_report(v)
    db.commit()
    audit(db, user.id, "plate_manual_set", f"violation={vid} plate={v.plate_text}")
    return {"id": v.id, "plate_text": v.plate_text, "plate_source": "manual"}

@router.post("/{vid}/finalize")
def finalize(vid: int, db: Session = Depends(get_db),
             user=Depends(require_position("officer", "administrator"))):
    v = _get(db, vid)
    if not v.plate_text:
        raise HTTPException(400, "Plate must be filled (OCR or manual) before finalizing")
    v.status = "finalized"
    v.report_path = generate_report(v)
    db.commit()
    audit(db, user.id, "violation_finalized", f"violation={vid}")
    return {"id": v.id, "status": v.status}

@router.post("/{vid}/reject")
def reject(vid: int, db: Session = Depends(get_db),
           user=Depends(require_position("administrator"))):   # evidence destruction = admin only
    v = _get(db, vid)
    for p in (v.image_path, v.report_path):
        if p and os.path.exists(p):
            os.remove(p)
    v.status = "deleted"
    db.commit()
    audit(db, user.id, "evidence_deleted", f"violation={vid}")
    return {"id": v.id, "status": v.status}

@router.get("/{vid}/report")
def download_report(vid: int, db: Session = Depends(get_db), 
                    user: User = Depends(get_current_user)):
    v = db.get(Violation, vid)
    if not v:
        raise HTTPException(404, "Violation not found")

    from ..services.templates import load_ui_settings, fill_template
    from fastapi.responses import FileResponse
    
    report_dir = Path(settings.STORAGE_ROOT) / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    report_path = report_dir / f"report_{v.id}.docx"

    # Check if a custom template is configured
    s = load_ui_settings()
    default_tpl = s.get("default_template")

    if default_tpl:
        # --- CUSTOM TEMPLATE PATH ---
        try:
            fill_template(default_tpl, {
                "plate_text": v.plate_text,
                "event_id": v.event_id,
                "violation_type": v.violation_type,
                "captured_at": v.captured_at,
                "status": v.status,
                "confidence": v.confidence,
                "officer_id": user.uid  # Stamp the downloading officer's ID
            }, str(report_path))
        except FileNotFoundError as e:
            raise HTTPException(404, str(e))
        except Exception as e:
            raise HTTPException(500, f"Template processing error: {e}")
    else:
        # --- LEGACY HARDCODED PATH ---
        from docx import Document
        doc = Document()
        doc.add_heading(f"ROADGUARD VIOLATION REPORT", 0)
        doc.add_paragraph(f"Event ID: {v.event_id}")
        doc.add_paragraph(f"Plate Number: {v.plate_text or 'NOT FILLED'}")
        doc.add_paragraph(f"Violation Type: {v.violation_type.replace('_', ' ').title()}")
        doc.add_paragraph(f"Date Captured: {v.captured_at}")
        doc.add_paragraph(f"Status: {v.status.upper()}")
        doc.add_paragraph(f"Confidence: {v.confidence}")
        doc.add_paragraph(f"\nReport generated by: {user.uid} on {utcnow()}")
        doc.save(str(report_path))

    return FileResponse(
        str(report_path), 
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"Violation_Report_{v.event_id}.docx"
    )

@router.get("/{vid}/image")
def download_image(vid: int, db: Session = Depends(get_db),
                   user=Depends(require_position(*ANY))):
    v = _get(db, vid)
    if not v.image_path or not os.path.exists(v.image_path):
        raise HTTPException(404, "Image not found")
    return FileResponse(v.image_path, filename=f"evidence_{v.id}.jpg")